from __future__ import annotations

import io
import zipfile

from docx import Document

from docgen.core.preflight import preflight


def _docx_with(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _inject_external_rel(docx_bytes: bytes, *, rel_type: str, target: str) -> bytes:
    """Rewrite word/_rels/document.xml.rels to include an external asset relationship."""
    rel_path = "word/_rels/document.xml.rels"
    with zipfile.ZipFile(io.BytesIO(docx_bytes), "r") as zin:
        parts = {name: zin.read(name) for name in zin.namelist()}
    rels_xml = parts.get(rel_path, b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                                    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
    rels_xml = rels_xml.replace(
        b"</Relationships>",
        (f'<Relationship Id="rIdEvil" Type="{rel_type}" '
         f'Target="{target}" TargetMode="External"/></Relationships>').encode(),
    )
    parts[rel_path] = rels_xml
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return out.getvalue()


def test_clean_docx_accepts(sample_docx_bytes: bytes) -> None:
    report = preflight(sample_docx_bytes)
    assert report.outcome == "accept"
    assert report.findings == []


def test_broken_bytes_rejected() -> None:
    report = preflight(b"not a docx at all")
    assert report.outcome == "reject"
    assert any(f.code == "opc.invalid" for f in report.findings)


def test_tracked_changes_rejected(tracked_change_docx_bytes: bytes) -> None:
    report = preflight(tracked_change_docx_bytes)
    assert report.outcome == "reject"
    assert any(f.code == "docx.tracked_changes" for f in report.findings)


def test_detects_jinja_variables_after_injection() -> None:
    doc = Document()
    doc.add_paragraph("Hello {{ member_name }}, your claim {{ claim_id }} is denied.")
    buf = io.BytesIO()
    doc.save(buf)

    report = preflight(buf.getvalue())
    assert report.outcome == "accept"
    assert set(report.detected_variables) == {"member_name", "claim_id"}


def test_empty_loop_no_guard_warns() -> None:
    report = preflight(_docx_with(["{% for o in orders %}{{ o }}{% endfor %}"]))
    assert report.outcome == "accept-with-warnings"
    assert any(f.code == "jinja.empty_loop_no_guard" for f in report.findings)


def test_guarded_loop_no_warning() -> None:
    report = preflight(
        _docx_with(["{% if orders %}{% for o in orders %}{{ o }}{% endfor %}{% endif %}"])
    )
    assert report.outcome == "accept"
    assert not any(f.code == "jinja.empty_loop_no_guard" for f in report.findings)


def test_for_else_no_warning() -> None:
    report = preflight(_docx_with(["{% for o in orders %}{{ o }}{% else %}(none){% endfor %}"]))
    assert report.outcome == "accept"
    assert not any(f.code == "jinja.empty_loop_no_guard" for f in report.findings)


def test_include_directive_rejected() -> None:
    report = preflight(_docx_with(["{% include 'other.docx' %}"]))
    assert report.outcome == "reject"
    assert any(
        f.code == "jinja.dangerous_construct" and "Include" in f.message for f in report.findings
    )


def test_dangerous_attribute_rejected() -> None:
    report = preflight(_docx_with(["{{ member.__class__ }}"]))
    assert report.outcome == "reject"
    assert any(
        f.code == "jinja.dangerous_construct" and "__class__" in f.message for f in report.findings
    )


def test_deep_attribute_chain_warns() -> None:
    report = preflight(_docx_with(["{{ member.profile.address.city }}"]))
    assert report.outcome == "accept-with-warnings"
    hits = [f for f in report.findings if f.code == "jinja.deep_attribute_chain"]
    assert len(hits) == 1
    assert "member.profile.address.city" in hits[0].message


def test_deep_chain_with_default_no_warning() -> None:
    report = preflight(_docx_with(["{{ member.profile.address.city | default('') }}"]))
    assert report.outcome == "accept"
    assert not any(f.code == "jinja.deep_attribute_chain" for f in report.findings)


def test_loop_var_shadowed_warns() -> None:
    report = preflight(
        _docx_with(["{% for x in xs %}{% for x in x.ys %}{{ x }}{% endfor %}{% endfor %}"])
    )
    assert report.outcome == "accept-with-warnings"
    assert any(f.code == "jinja.loop_var_shadowed" for f in report.findings)


def test_external_image_rejected() -> None:
    poisoned = _inject_external_rel(
        _docx_with(["hello"]),
        rel_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        target="http://169.254.169.254/latest/meta-data/",
    )
    report = preflight(poisoned)
    assert report.outcome == "reject"
    hits = [f for f in report.findings if f.code == "docx.external_asset"]
    assert len(hits) == 1
    assert "169.254.169.254" in hits[0].message
    assert "image" in hits[0].message


def test_external_ole_object_rejected() -> None:
    poisoned = _inject_external_rel(
        _docx_with(["hello"]),
        rel_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
        target="file:///etc/passwd",
    )
    report = preflight(poisoned)
    assert report.outcome == "reject"
    assert any(f.code == "docx.external_asset" for f in report.findings)


def test_external_hyperlink_allowed() -> None:
    """Hyperlinks only navigate on user click, not on render — should not trip the guard."""
    ok = _inject_external_rel(
        _docx_with(["hello"]),
        rel_type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        target="https://example.com",
    )
    report = preflight(ok)
    assert report.outcome == "accept"
    assert not any(f.code == "docx.external_asset" for f in report.findings)
