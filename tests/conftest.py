from __future__ import annotations

import io
import os
import uuid

import bcrypt
import pytest
import pytest_asyncio
from docx import Document

# Configure environment variables BEFORE importing app modules so that
# `get_settings()` picks them up on first access.
os.environ.setdefault(
    "DOCGEN_ADMIN_PASSWORD_HASH",
    bcrypt.hashpw(b"testpass", bcrypt.gensalt()).decode(),
)
os.environ.setdefault("DOCGEN_SESSION_SECRET", "test-session-secret")
os.environ.setdefault("DOCGEN_WEBHOOK_SIGNING_SECRET", "test-webhook-secret")

from sqlalchemy import select

from docgen.auth import generate_api_key
from docgen.store.db import async_engine, sync_session
from docgen.store.models import ApiKey, Tenant


@pytest.fixture(autouse=True)
def _dispose_async_engine_between_tests():
    """asyncpg's connection pool is bound to the event loop that opened each
    connection. `TestClient` spins up a fresh event loop per test; if the pool
    still holds handles from a previous loop, we get 'Event loop is closed' or
    'attached to a different loop'. Dispose the pool before each test so
    connections are always minted in the current loop.

    The redis async client used by RateLimitMiddleware has the same lifetime
    problem, so tear it down here too.
    """
    import asyncio as _a

    from docgen.ratelimit import close_redis

    async def _dispose():
        await async_engine.dispose()
        # Wipe login rate-limit state so tests hitting /ui/login don't share
        # a bucket. Since XFF is no longer trusted by default, every request
        # from TestClient resolves to peer "testclient", so distinct tests
        # would otherwise collide on the same key.
        from docgen.ratelimit import get_redis
        try:
            r = get_redis()
            keys = [k async for k in r.scan_iter(match="login_rl:*")]
            if keys:
                await r.delete(*keys)
        except Exception:
            pass
        await close_redis()

    try:
        _a.run(_dispose())
    except RuntimeError:
        pass
    yield

DEFAULT_TENANT_ID = uuid.UUID("00000000-0000-0000-0000-000000000001")


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Denial Letter", level=1)
    doc.add_paragraph("Dear [MEMBER_NAME],")
    doc.add_paragraph(
        "Your claim [CLAIM_ID] submitted on [CLAIM_DATE] has been denied. "
        "The amount in question was [AMOUNT] dollars."
    )
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph("<<agent_name>>")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def tracked_change_docx_bytes() -> bytes:
    """Docx with a tracked-change ins element in body.xml."""
    import zipfile

    doc = Document()
    doc.add_paragraph("Hello [NAME].")
    buf = io.BytesIO()
    doc.save(buf)
    base = buf.getvalue()

    with zipfile.ZipFile(io.BytesIO(base), "r") as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    xml = parts["word/document.xml"].decode("utf-8")
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    inject = (
        f'<w:ins xmlns:w="{W}" w:id="1" w:author="tester" w:date="2026-08-06T00:00:00Z">'
        f'<w:r><w:t xml:space="preserve"> tracked </w:t></w:r></w:ins>'
    )
    xml = xml.replace("</w:body>", inject + "</w:body>")
    parts["word/document.xml"] = xml.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, d in parts.items():
            zout.writestr(n, d)
    return out.getvalue()


@pytest.fixture
def default_tenant() -> uuid.UUID:
    with sync_session() as s:
        t = s.get(Tenant, DEFAULT_TENANT_ID)
        if t is None:
            t = Tenant(id=DEFAULT_TENANT_ID, name="default", ai_provider_config={})
            s.add(t)
            s.commit()
        return DEFAULT_TENANT_ID


@pytest.fixture
def api_key(default_tenant):
    raw, prefix, h = generate_api_key()
    with sync_session() as s:
        k = ApiKey(tenant_id=default_tenant, name="test", key_prefix=prefix, key_hash=h)
        s.add(k)
        s.commit()
    yield raw
    with sync_session() as s:
        result = s.execute(select(ApiKey).where(ApiKey.key_hash == h))
        obj = result.scalar_one_or_none()
        if obj:
            s.delete(obj)
            s.commit()


@pytest.fixture
def auth_headers(api_key):
    return {"Authorization": f"Bearer {api_key}"}
