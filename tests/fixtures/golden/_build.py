"""Regenerate golden template .docx files. Idempotent — run when adding a new golden.

Usage:  python tests/fixtures/golden/_build.py
Commit the resulting template.docx files.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

HERE = Path(__file__).parent


def build_denial_letter() -> None:
    d = Document()
    d.add_heading("Claim Denial Notice", level=1)
    d.add_paragraph("Date: {{ letter_date }}")
    d.add_paragraph("Member: {{ member_name }}")
    d.add_paragraph("Member ID: {{ member_id }}")
    d.add_paragraph(
        "Claim {{ claim_id }} submitted on {{ claim_date }} for US$ {{ amount }} "
        "has been denied."
    )
    d.add_paragraph("Reason: {{ denial_reason }}.")
    d.add_paragraph(
        "You may appeal within 60 days by calling {{ appeal_phone }} or writing to "
        "the address on the reverse of this notice."
    )
    d.add_paragraph("Sincerely,")
    d.add_paragraph("{{ agent_name }}, Claims Adjudicator")
    d.save(HERE / "denial_letter" / "template.docx")


def build_billing_statement() -> None:
    d = Document()
    d.add_heading("Monthly Statement", level=1)
    d.add_paragraph("Account: {{ account_number }}")
    d.add_paragraph("Statement date: {{ statement_date }}")
    d.add_paragraph("Line items follow.")

    # docxtpl table for-loop: {%tr for %} and {%tr endfor %} live in their own
    # rows (which vanish at render). The middle row is the one that repeats.
    tbl = d.add_table(rows=4, cols=3)
    hdr = tbl.rows[0].cells
    hdr[0].text = "Date"
    hdr[1].text = "Description"
    hdr[2].text = "Amount"
    tbl.rows[1].cells[0].text = "{%tr for line in lines %}"
    row = tbl.rows[2].cells
    row[0].text = "{{ line.date }}"
    row[1].text = "{{ line.description }}"
    row[2].text = "{{ line.amount }}"
    tbl.rows[3].cells[0].text = "{%tr endfor %}"

    d.add_paragraph("Total due: US$ {{ total_due }}")
    d.add_paragraph("Due by: {{ due_date }}")
    d.save(HERE / "billing_statement" / "template.docx")


if __name__ == "__main__":
    build_denial_letter()
    build_billing_statement()
    print("Rebuilt goldens in", HERE)
